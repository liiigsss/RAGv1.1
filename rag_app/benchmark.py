from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ROOT, settings


BENCHMARK_PDF = ROOT / "pdfs - Google Sheets.pdf"
JSONL_SOURCE = settings.source_dir / "LongDocURL_public.jsonl"

FIELDS = [
    "序号", "数据集", "doc_id", "question_id", "task_tag", "question_type", "answer_format",
    "问题", "标准答案", "证据页码", "证据类型", "对应原始PDF", "文档可用",
    "基准PDF页码", "基准PDF题目原文", "基准PDF答案原文", "生成答案", "检索引用",
    "语义正确(0/1)", "数值正确(0/1)", "引用正确(0/1)", "通过(公式)", "复核备注",
]


def _spaces(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _compact(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _answer_text(value: Any) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _raw_pdf_map() -> dict[str, Path]:
    result = {}
    for path in sorted(settings.cedar_dir.glob("*.pdf")):
        if "remediated" in path.stem.casefold():
            continue
        stem = path.stem.removeprefix("LDU__")
        result[stem] = path
    return result


def read_benchmark_rows() -> list[dict[str, Any]]:
    if not BENCHMARK_PDF.exists():
        raise FileNotFoundError(f"找不到测试基准 PDF: {BENCHMARK_PDF}")
    if not JSONL_SOURCE.exists():
        raise FileNotFoundError(f"找不到 LongDocURL 来源数据: {JSONL_SOURCE}")
    source_rows = [json.loads(line) for line in JSONL_SOURCE.read_text(encoding="utf-8").splitlines() if line.strip()]
    by_question_id = {row["question_id"]: row for row in source_rows}
    raw_pdfs = _raw_pdf_map()
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    with pdfplumber.open(BENCHMARK_PDF) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables() or []
            for table in tables:
                for row in table:
                    if len(row) <= 10:
                        continue
                    question_id = _compact(row[6])
                    if not question_id or question_id == "question_id":
                        continue
                    if question_id in seen:
                        continue
                    seen.add(question_id)
                    doc_id = _compact(row[0])
                    source_label = _spaces(row[1])
                    if question_id in by_question_id:
                        canonical = by_question_id[question_id]
                        doc_no = str(canonical["doc_no"])
                        dataset = "LongDocURL"
                        question = canonical["question"]
                        answer = _answer_text(canonical["answer"])
                        task_tag = canonical.get("task_tag", "")
                        question_type = canonical.get("question_type", "")
                        answer_format = canonical.get("answer_format", "")
                        evidence_pages = canonical.get("evidence_pages", [])
                        evidence_sources = canonical.get("evidence_sources", [])
                        if not doc_id:
                            doc_id = f"LDU__{doc_no}"
                    else:
                        if source_label.casefold().startswith("finance"):
                            dataset = "FinanceBench"
                        else:
                            raise ValueError(f"基准 PDF 中的 question_id 未能匹配来源 JSONL：{question_id}")
                        doc_no = doc_id.removeprefix("FB__")
                        question = _spaces(row[9])
                        answer = _spaces(row[10])
                        task_tag = _spaces(row[2])
                        question_type = _spaces(row[7])
                        answer_format = _spaces(row[8])
                        evidence_pages = []
                        evidence_sources = []
                    pdf_path = raw_pdfs.get(doc_no)
                    result.append({
                        "序号": len(result) + 1,
                        "数据集": dataset,
                        "doc_id": doc_id,
                        "question_id": question_id,
                        "task_tag": task_tag,
                        "question_type": question_type,
                        "answer_format": answer_format,
                        "问题": question,
                        "标准答案": answer,
                        "证据页码": ", ".join(map(str, evidence_pages)) if evidence_pages else "",
                        "证据类型": ", ".join(map(str, evidence_sources)) if evidence_sources else "",
                        "对应原始PDF": pdf_path.name if pdf_path else "",
                        "文档可用": "是" if pdf_path else "否",
                        "基准PDF页码": page_number,
                        "基准PDF题目原文": _spaces(row[9]),
                        "基准PDF答案原文": _spaces(row[10]),
                        "生成答案": "",
                        "检索引用": "",
                        "语义正确(0/1)": "",
                        "数值正确(0/1)": "",
                        "引用正确(0/1)": "",
                        "通过(公式)": "",
                        "复核备注": "" if pdf_path else "基准题保留，但 Cedar 中缺少对应原始 PDF。",
                    })
    return result


def build_benchmark() -> tuple[Path, Path, list[dict[str, Any]]]:
    settings.benchmark_dir.mkdir(parents=True, exist_ok=True)
    rows = read_benchmark_rows()
    jsonl_path = settings.benchmark_dir / "benchmark.jsonl"
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    xlsx_path = settings.benchmark_dir / "benchmark_master.xlsx"
    write_workbook(rows, xlsx_path, include_formulas=False)
    return jsonl_path, xlsx_path, rows


def write_workbook(rows: list[dict[str, Any]], path: Path, include_formulas: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试基准与结果"
    sheet.append(FIELDS)
    for row_number, row in enumerate(rows, start=2):
        values = [row.get(field, "") for field in FIELDS]
        if include_formulas:
            values[FIELDS.index("通过(公式)")] = (
                f'=IF(OR(S{row_number}="",T{row_number}="",U{row_number}=""),"",'
                f'IF(AND(S{row_number}=1,T{row_number}=1,U{row_number}=1),1,0))'
            )
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 16, 20, 38, 18, 18, 14, 54, 44, 16, 18, 28, 12, 12, 50, 44, 50, 40, 16, 16, 16, 14, 44]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234A6F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary = workbook.create_sheet("准确率汇总")
    summary.append(["指标", "数值", "说明"])
    last_row = max(2, len(rows) + 1)
    summary.append(["基准题数", f"=COUNTA('测试基准与结果'!D2:D{last_row})", "与测试基准逐题对应，包含文档缺失题目"])
    summary.append(["文档可用题数", f'=COUNTIF(\'测试基准与结果\'!M2:M{last_row},"是")', "可进入本地检索与生成流程"])
    summary.append(["已评分题数", f'=COUNT(\'测试基准与结果\'!V2:V{last_row})', "语义、数值、引用三项均完成后计入"])
    summary.append(["准确率", f'=IFERROR(AVERAGE(\'测试基准与结果\'!V2:V{last_row}),"")', "通过题数 / 已评分题数；人工评分列可编辑"])
    summary.append(["原始基准 PDF SHA256", hashlib.sha256(BENCHMARK_PDF.read_bytes()).hexdigest(), "用于确认基准版本"])
    summary.append(["LongDocURL JSONL SHA256", hashlib.sha256(JSONL_SOURCE.read_bytes()).hexdigest(), "用于确认 canonical 问题与答案来源"])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 72
    summary.column_dimensions["C"].width = 62
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234A6F")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide = workbook.create_sheet("字段说明")
    guide.append(["评分字段", "填写规则"])
    guide.append(["语义正确(0/1)", "模型评审预填，人工可修改；答案与标准答案含义一致填 1，否则填 0。"])
    guide.append(["数值正确(0/1)", "关键金额、数量、期间、单位、币种和正负号均正确才填 1；无数值要求的题目填 1。"])
    guide.append(["引用正确(0/1)", "引用的文档和页码能在检索证据中找到，且证据支持答案，填 1。"])
    guide.append(["通过(公式)", "三项评分均为 1 时自动计 1；任一项为 0 时计 0；有评分空缺时保持空白。"])
    guide.append(["基准PDF原文", "保留从 Google Sheets 导出 PDF 解析出的题目与答案；LongDocURL 的标准文本从原始 JSONL 按 question_id 对齐，以修复 PDF 换行导致的断词。"])
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234A6F")
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def load_benchmark() -> list[dict[str, Any]]:
    path = settings.benchmark_dir / "benchmark.jsonl"
    if not path.exists():
        _, _, rows = build_benchmark()
        return rows
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_result_workbook(rows: list[dict[str, Any]], path: Path) -> None:
    write_workbook(rows, path, include_formulas=True)
