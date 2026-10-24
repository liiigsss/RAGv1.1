from __future__ import annotations

import json
import os
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .benchmark import load_benchmark
from .config import settings


def _find_key(row: dict[str, Any], prefix: str) -> str | None:
    return next((key for key in row if key.startswith(prefix)), None)


def _get(row: dict[str, Any], prefix: str, default: Any = None) -> Any:
    key = _find_key(row, prefix)
    return row.get(key, default) if key else default


def _is_one(value: Any) -> bool:
    return value == 1 or str(value).strip() == "1"


def _is_pass_or_na(value: Any) -> bool:
    return _is_one(value) or str(value or "").strip().upper() == "N/A"


def _status(row: dict[str, Any]) -> str:
    return str(_get(row, "评测状态", "未评测") or "未评测")


def _strict_pass(row: dict[str, Any]) -> bool:
    return (
        _status(row) == "完成"
        and _is_one(_get(row, "语义正确"))
        and _is_pass_or_na(_get(row, "数值正确"))
        and _is_pass_or_na(_get(row, "List条目匹配"))
        and _is_pass_or_na(_get(row, "格式符合"))
        and _is_one(_get(row, "引用是否命中检索证据"))
        and _is_one(_get(row, "证据内容支持"))
    )


def _rate(rows: list[dict[str, Any]], prefix: str) -> float | None:
    values = [_get(row, prefix) for row in rows]
    values = [value for value in values if value not in (None, "") and str(value).strip().upper() != "N/A"]
    if not values:
        return None
    return sum(_is_one(value) for value in values) / len(values)


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        return "".join(char if char in "\t\n\r" or ord(char) >= 0x20 else " " for char in value)
    return value


def _style_sheet(sheet, widths: list[int] | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    if sheet.max_row > 1 and sheet.max_column > 0:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="234A6F")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, sheet.max_column + 1):
        if widths and column <= len(widths):
            width = widths[column - 1]
        else:
            header = str(sheet.cell(1, column).value or "")
            width = min(max(len(header) + 4, 14), 42)
        sheet.column_dimensions[get_column_letter(column)].width = width


def _read_latest_checkpoint(path: Path) -> dict[str, dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    if not path.exists():
        return latest
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            row = json.loads(line)
            question_id = row.get("question_id")
            if question_id:
                latest[str(question_id)] = row
    return latest


def _flatten_details(row: dict[str, Any], prefix: str) -> list[dict[str, Any]]:
    key = _find_key(row, prefix)
    items = row.get(key, []) if key else []
    if not isinstance(items, list):
        return []
    return [
        {"序": row.get("序", ""), "question_id": row.get("question_id", ""), **item}
        for item in items
        if isinstance(item, dict)
    ]


def _write_table(sheet, headers: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])
    _style_sheet(sheet)


def _build_dataset_report(
    dataset: str,
    benchmark_rows: list[dict[str, Any]],
    latest: dict[str, dict[str, Any]],
    path: Path,
) -> Path:
    rows = []
    for source in benchmark_rows:
        row = dict(source)
        row.update(latest.get(str(source.get("question_id", "")), {}))
        rows.append(row)
    def sort_key(row: dict[str, Any]) -> tuple[int, str]:
        try:
            sequence = int(row.get("序", 0))
        except (TypeError, ValueError):
            sequence = 0
        return sequence, str(row.get("question_id", ""))

    rows.sort(key=sort_key)
    for row in rows:
        row["严格通过(0/1)"] = 1 if _strict_pass(row) else (0 if _status(row) == "完成" else "")

    nested_prefixes = ("检索候选明细", "发送证据明细")
    result_headers = list(OrderedDict.fromkeys(
        key for row in rows for key in row if not any(key.startswith(prefix) for prefix in nested_prefixes)
    ))
    failed_rows = [row for row in rows if not _strict_pass(row)]
    retrieval_rows = [item for row in failed_rows for item in _flatten_details(row, "检索候选明细")]
    evidence_rows = [item for row in failed_rows for item in _flatten_details(row, "发送证据明细")]

    completed = [row for row in rows if _status(row) == "完成"]
    missing_pdf = [row for row in rows if _status(row) == "缺少原始PDF"]
    unprocessed = len(rows) - len(completed) - len(missing_pdf)
    passed = sum(_strict_pass(row) for row in rows)
    summary_rows = [
        ("数据集", dataset, "评测输入数据集。"),
        ("基准问题数", len(rows), "从 benchmark.jsonl 读取。"),
        ("已完成评测数", len(completed), "评测状态为“完成”的题目。"),
        ("缺少原始 PDF", len(missing_pdf), "不计入已完成题目的准确率分母。"),
        ("未完成/未评测", unprocessed, "可用同一 checkpoint 续跑。"),
        ("严格通过数", passed, "语义、适用的数值/List/格式项、引用命中及证据支持均通过。"),
        ("严格准确率", passed / len(completed) if completed else None, "严格通过数 / 已完成评测数。"),
        ("语义正确率", _rate(completed, "语义正确"), "语义正确项比例。"),
        ("数值正确率", _rate(completed, "数值正确"), "排除 N/A 后的数值正确项比例。"),
        ("List 条目匹配率", _rate(completed, "List条目匹配"), "排除 N/A 后的 List 匹配比例。"),
        ("格式符合率", _rate(completed, "格式符合"), "排除 N/A 后的格式符合比例。"),
        ("引用命中率", _rate(completed, "引用是否命中检索证据"), "引用是否命中实际检索证据。"),
        ("证据支持率", _rate(completed, "证据内容支持"), "实际发送的证据内容是否支持回答。"),
        ("生成模型", next((row.get("生成模型") for row in completed if row.get("生成模型")), ""), "来自评测记录，不包含 API 密钥。"),
        ("Embedding 模型", next((row.get("Embedding模型") for row in completed if row.get("Embedding模型")), ""), "来自评测记录。"),
        ("Reranker 模型", next((row.get("Reranker模型") for row in completed if row.get("Reranker模型")), ""), "来自评测记录。"),
    ]

    workbook = Workbook()
    summary = workbook.active
    summary.title = "准确率概览"
    summary.append(["指标", "值", "说明"])
    for label, value, note in summary_rows:
        summary.append([label, value, note])
    _style_sheet(summary, [28, 42, 82])
    for row_index in range(2, summary.max_row + 1):
        if str(summary.cell(row_index, 1).value).endswith("率"):
            summary.cell(row_index, 2).number_format = "0.0%"

    results = workbook.create_sheet("评测结果")
    _write_table(results, result_headers, rows)
    retrieval = workbook.create_sheet("检索明细")
    retrieval_headers = list(OrderedDict.fromkeys(key for row in retrieval_rows for key in row))
    _write_table(retrieval, retrieval_headers or ["序", "question_id"], retrieval_rows)
    evidence = workbook.create_sheet("发送证据")
    evidence_headers = list(OrderedDict.fromkeys(key for row in evidence_rows for key in row))
    _write_table(evidence, evidence_headers or ["序", "question_id"], evidence_rows)
    guide = workbook.create_sheet("字段说明")
    guide.append(["字段", "说明"])
    guide.append(["严格通过", "评测完成，语义正确，适用的数值/List/格式项通过，引用命中检索证据且证据内容支持回答。"])
    guide.append(["准确率分母", "已完成评测题目；缺少原始 PDF 的题目单独统计，不计入分母。"])
    guide.append(["检索明细/发送证据", "仅列出未严格通过题目的检索候选和实际发送给生成模型的证据。"])
    guide.append(["续跑记录", "原始 JSONL checkpoint 保留在 outputs 中；工作簿可从记录重新生成。"])
    _style_sheet(guide, [30, 110])

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_evaluation_reports(checkpoint_path: Path) -> list[Path]:
    latest = _read_latest_checkpoint(checkpoint_path)
    by_dataset: dict[str, list[dict[str, Any]]] = OrderedDict()
    for row in load_benchmark():
        dataset = str(row.get("数据集", "未分类"))
        by_dataset.setdefault(dataset, []).append(row)

    version = os.getenv("RAG_EVALUATION_VERSION", "v1.1")
    reports = []
    datasets_with_records = {
        str(row.get("数据集", "未分类")) for row in latest.values()
    }
    for dataset, rows in by_dataset.items():
        if datasets_with_records and dataset not in datasets_with_records:
            continue
        safe_dataset = "".join(char if char.isalnum() or char in "-_" else "_" for char in dataset)
        report_path = settings.output_dir / f"{safe_dataset}_accuracy_{version}.xlsx"
        reports.append(_build_dataset_report(dataset, rows, latest, report_path))
    return reports
